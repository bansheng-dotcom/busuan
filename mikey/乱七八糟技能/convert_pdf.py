# -*- coding: utf-8 -*-
"""Convert 自然流·救赎之道 PDF -> TXT + DOCX."""
import re
import pymupdf
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

PDF = r"D:/cc/自然流·救赎之道_微信读书版.pdf"
TXT = r"D:/cc/自然流·救赎之道_微信读书版.txt"
DOCX = r"D:/cc/自然流·救赎之道_微信读书版.docx"

SENTENCE_END = set("。！？!?；;…")
QUOTE_END = set("”’」』）》】〉\"'")

CHAPTER_RE = re.compile(r"^\s*第[0-9一二三四五六七八九十百零两]+[章卷节篇]\s*")


def extract_paragraphs():
    doc = pymupdf.open(PDF)
    paragraphs = []
    cur = []

    def flush():
        nonlocal cur
        if cur:
            paragraphs.append("".join(cur).strip())
            cur = []

    for page in doc:
        text = page.get_text("text")
        for raw in text.split("\n"):
            line = raw.strip()
            if not line:
                flush()
                continue
            # chapter heading on its own -> own paragraph
            if CHAPTER_RE.match(line) and len(line) <= 40:
                flush()
                paragraphs.append(line)
                continue
            cur.append(line)
            if line[-1] in SENTENCE_END or line[-1] in QUOTE_END:
                flush()
        flush()

    # drop fully-empty, collapse blank paragraphs already handled
    return [p for p in paragraphs if p]


def build_txt(paragraphs):
    with open(TXT, "w", encoding="utf-8-sig", newline="\n") as f:
        f.write("\n\n".join(paragraphs))


def set_font(run, name_ascii, name_east, size_pt=None, bold=False):
    run.font.name = name_ascii
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name_ascii)
    rFonts.set(qn("w:hAnsi"), name_ascii)
    rFonts.set(qn("w:eastAsia"), name_east)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    run.font.bold = bold


def build_docx(paragraphs):
    doc = Document()

    # base style
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")

    for p in paragraphs:
        if CHAPTER_RE.match(p) and len(p) <= 40:
            h = doc.add_heading(level=1)
            run = h.add_run(p)
            set_font(run, "Microsoft YaHei", "微软雅黑", size_pt=16, bold=True)
        else:
            para = doc.add_paragraph()
            run = para.add_run(p)
            set_font(run, "Times New Roman", "宋体", size_pt=11)

    doc.save(DOCX)


def main():
    paragraphs = extract_paragraphs()
    print(f"paragraphs={len(paragraphs)}")
    build_txt(paragraphs)
    build_docx(paragraphs)
    print("done")


if __name__ == "__main__":
    main()
