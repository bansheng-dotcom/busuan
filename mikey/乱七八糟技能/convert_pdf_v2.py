# -*- coding: utf-8 -*-
"""Rebuild 自然流·救赎之道 DOCX with: chapter pagination, auto TOC,
bookmarks, cleaned blank lines, fonts/spacing. Then convert to .doc via Word."""
import re
import os
import pymupdf
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PDF = r"D:/cc/自然流·救赎之道_微信读书版.pdf"
DOCX = r"D:/cc/自然流·救赎之道_微信读书版.docx"
DOC = r"D:/cc/自然流·救赎之道_微信读书版.doc"

SENTENCE_END = set("。！？!?；;…")
QUOTE_END = set("”’」』）》】〉\"'")
CHAPTER_RE = re.compile(r"^\s*第([0-9]+)章\s*")
HEADING_RE = re.compile(r"^第([0-9]+)章[：: ]*(.+)")


# ---------------- extract ----------------
def extract_paragraphs():
    doc = pymupdf.open(PDF)
    paras, cur = [], []

    def flush():
        nonlocal cur
        if cur:
            paras.append("".join(cur).strip())
            cur = []

    for page in doc:
        for raw in page.get_text("text").split("\n"):
            line = raw.strip()
            if not line:
                flush()
                continue
            if CHAPTER_RE.match(line) and len(line) <= 40:
                flush()
                paras.append(line)
                continue
            cur.append(line)
            if line[-1] in SENTENCE_END or line[-1] in QUOTE_END:
                flush()
        flush()
    return [p for p in paras if p]


def find_toc(paras):
    idx = next((i for i, p in enumerate(paras) if p.strip() == "目 录"), None)
    canonical, toc_end = {}, idx
    if idx is not None:
        j = idx + 1
        while j < len(paras):
            m = CHAPTER_RE.match(paras[j])
            if not m:
                break
            num = int(m.group(1))
            if num in canonical:  # body starts repeating chapter numbers
                break
            canonical[num] = paras[j]
            toc_end = j
            j += 1
    return idx, toc_end, canonical


# ---------------- docx helpers ----------------
def style_font(style, east, ascii_font, size, bold=False):
    style.font.name = ascii_font
    style.font.size = Pt(size)
    style.font.bold = bold
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)


def run_font(run, east, ascii_font, size, bold=False, color=None):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)


def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-1" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "（目录占位：在 Word 中右键 → 更新域 即可生成）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r = run._r
    for el in (fld_begin, instr, fld_sep, t, fld_end):
        r.append(el)


def centered(text, east, ascii_font, size, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(p.add_run(text), east, ascii_font, size, bold)
    return p


# ---------------- build ----------------
paragraphs = extract_paragraphs()
toc_idx, toc_end, canonical = find_toc(paragraphs)
print(f"paragraphs={len(paragraphs)} chapters={len(canonical)} toc=[{toc_idx},{toc_end}]")

doc = Document()

# base body style: 宋体 小四(12pt) 1.5x
normal = doc.styles["Normal"]
style_font(normal, "宋体", "Times New Roman", 12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)

# Heading 1 style: 微软雅黑 三号(16pt) 加粗, 每章另起一页
h1 = doc.styles["Heading 1"]
style_font(h1, "微软雅黑", "Microsoft YaHei", 16, bold=True)
h1.font.color.rgb = RGBColor(0, 0, 0)
h1.paragraph_format.page_break_before = True
h1.paragraph_format.space_before = Pt(6)
h1.paragraph_format.space_after = Pt(12)
h1.paragraph_format.line_spacing = 1.3

# ---- cover / title page ----
centered("自然流 · 救赎之道", "微软雅黑", "Microsoft YaHei", 26, bold=True)
centered("从三伪神到全流程案例集 —— 完整教学体系", "微软雅黑", "Microsoft YaHei", 15)
centered("Mike · 神圣教学", "宋体", "Times New Roman", 12)
sp = centered("", "宋体", "Times New Roman", 12)
sp.paragraph_format.space_after = Pt(24)

# 版权声明 / 全书共70章 (paragraphs[1], paragraphs[2])
for p in paragraphs[1:3]:
    doc.add_paragraph(p)

# ---- auto TOC ----
toc_head = centered("目  录", "微软雅黑", "Microsoft YaHei", 18, bold=True)
toc_head.paragraph_format.space_before = Pt(18)
add_toc_field(doc.add_paragraph())

# ---- body (skip cover[0] + 版权[1,2] + manual toc[3..toc_end]) ----
body_start = max(toc_end + 1, 3)
for p in paragraphs[body_start:]:
    m = HEADING_RE.match(p)
    title = None
    if m:
        num = int(m.group(1))
        if num in canonical:
            can = canonical[num]
            if can.startswith(p) or p == can:
                title = can
    if title:
        doc.add_heading(title, level=1)
    else:
        doc.add_paragraph(p)

doc.save(DOCX)
print("docx saved")


# ---------------- Word COM: update TOC, add bookmarks, save .doc ----------------
import win32com.client
import pythoncom

pythoncom.CoInitialize()
word = win32com.client.DispatchEx("Word.Application")
word.Visible = False
word.DisplayAlerts = 0
w = word.Documents.Open(os.path.abspath(DOCX))

# update all fields (TOC)
try:
    if w.TablesOfContents.Count > 0:
        w.TablesOfContents(1).Update()
except Exception as e:
    print("toc update warn:", e)
try:
    for story in w.StoryRanges:
        story.Fields.Update()
except Exception as e:
    print("field update warn:", e)

# bookmarks for each chapter heading
added = 0
for p in w.Paragraphs:
    name = p.Style.NameLocal or ""
    if ("标题" in name and "1" in name) or "Heading 1" in name:
        txt = p.Range.Text.strip()
        m = re.match(r"第([0-9]+)章", txt)
        bm = f"第{m.group(1)}章" if m else f"Ch{added + 1}"
        try:
            w.Bookmarks.Add(bm, p.Range)
            added += 1
        except Exception:
            pass
print("bookmarks added:", added)

w.SaveAs2(os.path.abspath(DOCX), FileFormat=16)  # docx
w.SaveAs2(os.path.abspath(DOC), FileFormat=0)    # legacy doc
w.Close(False)
word.Quit()
print("done")
